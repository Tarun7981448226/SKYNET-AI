from app.rendering.resume_render import generate_docx, generate_pdf
from app.resume.schema import Contact
from app.tailoring.schema import TailoredItem, TailoredResumeContent


def _tailored() -> TailoredResumeContent:
    return TailoredResumeContent(
        job_id=1,
        contact=Contact(
            name="Jane Doe",
            email="jane@example.com",
            phone="+1-213-000-0000",
            location="Los Angeles, CA",
            links=["www.linkedin.com/in/janedoe", "https://github.com/janedoe"],
        ),
        summary=[TailoredItem(id="summary-main", section="summary", bullets=["MSCS student."])],
        education=[
            TailoredItem(id="edu-usc", section="education", dates="2024 – 2026", bullets=["M.S. CS, USC."])
        ],
        experience=[
            TailoredItem(
                id="exp-one",
                section="experience",
                dates="2024",
                bullets=["Built a data pipeline.", "Reduced latency 30%."],
            )
        ],
        projects=[
            TailoredItem(id="project-a", section="projects", dates="2024", bullets=["Trained a model."])
        ],
        skills=[TailoredItem(id="skills-ml", section="skills", bullets=["ML: PyTorch"])],
    )


def test_generate_pdf_creates_single_page_readable_file(tmp_path):
    from PyPDF2 import PdfReader

    output_path = tmp_path / "resume.pdf"
    generate_pdf(_tailored(), str(output_path))

    assert output_path.exists()
    reader = PdfReader(str(output_path))
    assert len(reader.pages) == 1

    text = reader.pages[0].extract_text()
    assert "Jane Doe" in text
    assert "jane@example.com" in text
    assert "Built a data pipeline" in text
    assert "Trained a model" in text


def test_generate_pdf_has_real_hyperlinks(tmp_path):
    from PyPDF2 import PdfReader

    output_path = tmp_path / "resume.pdf"
    generate_pdf(_tailored(), str(output_path))

    reader = PdfReader(str(output_path))
    annotations = reader.pages[0].get("/Annots") or []
    uris = []
    for annot in annotations:
        obj = annot.get_object()
        action = obj.get("/A")
        if action and "/URI" in action:
            uris.append(action["/URI"])

    assert any(uri.startswith("mailto:jane@example.com") for uri in uris)
    assert any(uri.startswith("tel:") for uri in uris)
    assert any("linkedin.com" in uri for uri in uris)
    assert any("github.com" in uri for uri in uris)


def test_generate_docx_creates_readable_file_with_hyperlinks(tmp_path):
    from docx import Document

    output_path = tmp_path / "resume.docx"
    generate_docx(_tailored(), str(output_path))

    assert output_path.exists()
    doc = Document(str(output_path))
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "Jane Doe" in full_text
    assert "Built a data pipeline" in full_text

    # Confirm a real w:hyperlink element exists (not just plain text) for email.
    assert any("w:hyperlink" in p._p.xml for p in doc.paragraphs)
